from __future__ import annotations

from openpyxl import load_workbook

from dks_model.engine import run_model
from dks_model.io import load_inputs
from dks_model.workbook import SHEET_ORDER, export_workbook


def test_workbook_contains_required_tabs() -> None:
    inputs = load_inputs("inputs")
    outputs = run_model(inputs, scenario="base")
    output_path = "outputs/test_DKS_model.xlsx"

    export_workbook(outputs, inputs, output_path)

    workbook = load_workbook(output_path, read_only=True, data_only=False)
    assert workbook.sheetnames == SHEET_ORDER
    assert workbook["Model Audit & Quality Checks"]["B2"].value == "PASS"
    assert workbook["Probabilistic Valuation"]["A3"].value == "status"
    assert workbook["Thesis Catalysts Monitoring"]["A3"].value == "status"
    assert workbook["Forecast Variance Analysis"]["A3"].value == "status"
    assert workbook["Alternative Data & Backtesting"]["A3"].value == "Reserved for future alternative-data collection and validation."
